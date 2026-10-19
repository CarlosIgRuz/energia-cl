"""Exploración (una vez) de los ZIP diarios de "Costo Marginal Real" del sitio del Coordinador.

Recorre año → mes → día para algunos días de muestra, descarga el ZIP del día y registra en
data/logs/explorar_cmg_sitio.json: URLs, tamaño, archivos internos y sus primeras líneas (o, si son Excel, hojas y
primeras filas). No guarda los ZIP. ~3 requests por día de muestra, con pausa.
"""
from __future__ import annotations

import io
import json
import re
import sys
import time
import zipfile
from html import unescape
from pathlib import Path
from urllib.parse import urljoin

import requests

from energia.ingest.comun import LOGS, SITIO_CEN, USER_AGENT_NAVEGADOR, ahora_utc

BASE = f"{SITIO_CEN}/mercados/documentos/transferencias-economicas/costo-marginal-real/"
ANIOS = {
    "2026": BASE + "2026-costo-marginal-real/",
    "2025": BASE + "2025-costo-marginal-real/",
    "2024": BASE + "2024-costo-marginal-real-transferencias-economicas/",
}
_HREF = re.compile(r'href=["\']([^"\']+)["\']', re.IGNORECASE)


def enlaces(html: str, base: str) -> list[str]:
    return list(dict.fromkeys(urljoin(base, unescape(h)) for h in _HREF.findall(html)))


def muestra_archivo(nombre: str, datos: bytes) -> dict:
    info: dict = {"nombre": nombre, "bytes": len(datos)}
    low = nombre.lower()
    if low.endswith((".csv", ".txt", ".tsv")):
        texto = datos[:4000].decode("latin-1", errors="replace")
        info["primeras_lineas"] = texto.splitlines()[:8]
        info["n_lineas_aprox"] = datos.count(b"\n")
    elif low.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
            info["hojas"] = {}
            for ws in wb.worksheets[:6]:
                filas = []
                for i, fila in enumerate(ws.iter_rows(values_only=True)):
                    filas.append([None if v is None else str(v)[:40] for v in fila[:14]])
                    if i >= 9:
                        break
                info["hojas"][ws.title] = {"dimensiones": ws.max_row and f"{ws.max_row}x{ws.max_column}", "filas": filas}
        except Exception as e:  # noqa: BLE001 — es exploración: registrar y seguir
            info["error_excel"] = str(e)[:200]
    return info


def main(logs_dir: Path = LOGS, pausa_s: float = 3.0) -> int:
    s = requests.Session()
    s.headers.update({"User-Agent": USER_AGENT_NAVEGADOR})
    salida: dict = {"ejecutado_en": ahora_utc(), "muestras": {}, "requests": 0}

    def get(url: str) -> requests.Response:
        salida["requests"] += 1
        r = s.get(url, timeout=120)
        time.sleep(pausa_s)
        return r

    for anio, url_anio in ANIOS.items():
        m: dict = {"url_anio": url_anio}
        salida["muestras"][anio] = m
        try:
            r = get(url_anio)
            m["status_anio"] = r.status_code
            meses = [u for u in enlaces(r.text, url_anio) if "costo-marginal-real" in u and u.rstrip("/") != url_anio.rstrip("/")
                     and re.search(r"/(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)-", u)]
            m["meses"] = meses[:14]
            if not meses:
                continue
            url_mes = meses[0] if anio == "2026" else meses[len(meses) // 2]
            r = get(url_mes)
            m["url_mes"], m["status_mes"] = url_mes, r.status_code
            dias = [u for u in enlaces(r.text, url_mes) if re.search(r"/\d{2}-[a-z]+-\d{4}-costo-marginal-real", u)]
            m["n_dias"] = len(dias)
            if not dias:
                m["enlaces_mes"] = enlaces(r.text, url_mes)[:80]
                continue
            url_dia = dias[min(1, len(dias) - 1)]
            r = get(url_dia)
            m["url_dia"], m["status_dia"] = url_dia, r.status_code
            archivos = [u for u in enlaces(r.text, url_dia) if re.search(r"\.(zip|xlsx|csv|rar|7z)$", u, re.I)]
            m["archivos_dia"] = archivos
            for u in archivos[:2]:
                rz = get(u)
                arch = {"url": u, "status": rz.status_code, "bytes": len(rz.content)}
                if rz.status_code == 200 and u.lower().endswith(".zip"):
                    z = zipfile.ZipFile(io.BytesIO(rz.content))
                    arch["miembros"] = [{"nombre": i.filename, "bytes": i.file_size} for i in z.infolist()]
                    arch["muestras"] = [muestra_archivo(i.filename, z.read(i)) for i in z.infolist()[:6]
                                        if not i.is_dir()]
                m.setdefault("descargas", []).append(arch)
        except Exception as e:  # noqa: BLE001
            m["error"] = str(e)[:300]
        print(anio, json.dumps({k: v for k, v in m.items() if k in ("status_anio", "n_dias", "archivos_dia", "error")}))
    logs_dir.mkdir(parents=True, exist_ok=True)
    (logs_dir / "explorar_cmg_sitio.json").write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
